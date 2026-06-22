"""统一文档解析引擎 — 将各类文档解析为结构化内容供搜索和AI消费"""
import os
from dataclasses import dataclass, field
from datetime import datetime


@dataclass
class Section:
    title: str = ""
    content: str = ""
    subsections: list = field(default_factory=list)
    level: int = 1


@dataclass
class Table:
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    caption: str = ""


@dataclass
class ParsedDocument:
    filepath: str
    filename: str
    file_type: str
    text_content: str = ""
    structured_content: str = ""
    metadata: dict = field(default_factory=dict)
    sections: list[Section] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    summary: str = ""


def parse_document(filepath: str) -> ParsedDocument:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xlsx":
        return parse_xlsx(filepath)
    if ext == ".xls":
        return parse_xls(filepath)
    if ext == ".pptx":
        return parse_pptx(filepath)
    if ext == ".docx":
        return parse_docx(filepath)
    if ext == ".pdf":
        return parse_pdf(filepath)
    return ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type=ext.lstrip("."),
    )


def _build_metadata(filepath: str) -> dict:
    stat = os.stat(filepath)
    return {
        "file_size": stat.st_size,
        "modified_time": datetime.fromtimestamp(stat.st_mtime).isoformat(),
        "created_time": datetime.fromtimestamp(stat.st_ctime).isoformat(),
    }


# ── .docx 解析 ──────────────────────────────────────────


def parse_docx(filepath: str) -> ParsedDocument:
    from docx import Document

    doc = Document(filepath)
    docx_meta = doc.core_properties
    ext_meta = {
        "author": docx_meta.author or "",
        "title": docx_meta.title or "",
        "created": str(docx_meta.created or ""),
        "modified": str(docx_meta.modified or ""),
    }
    ext_meta.update(_build_metadata(filepath))

    parsed = ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type="docx",
        metadata=ext_meta,
    )

    text_parts = []
    md_parts = [f"# {parsed.filename}\n"]

    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        text_parts.append(t)
        if para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading", "").strip()
            try:
                lvl = min(int(level), 4)
            except ValueError:
                lvl = 1
            md_parts.append(f"\n{'#' * lvl} {t}\n")
        else:
            md_parts.append(t)

    for table in doc.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        rows = []
        for row in table.rows[1:]:
            rows.append([cell.text.strip() for cell in row.cells])
        parsed.tables.append(Table(headers=headers, rows=rows))

        md_parts.append("\n| " + " | ".join(headers) + " |")
        md_parts.append("|" + "|".join(["---"] * len(headers)) + "|")
        for row in rows:
            md_parts.append("| " + " | ".join(row) + " |")

    parsed.text_content = "\n".join(text_parts)
    parsed.structured_content = "\n".join(md_parts)
    parsed.summary = _summarize_text(text_parts)
    return parsed


# ── .pdf 解析 ───────────────────────────────────────────


def parse_pdf(filepath: str) -> ParsedDocument:
    import fitz

    pdf_doc = fitz.open(filepath)
    ext_meta = {
        "page_count": len(pdf_doc),
        "title": pdf_doc.metadata.get("title", ""),
        "author": pdf_doc.metadata.get("author", ""),
    }
    pdf_meta = _build_metadata(filepath)
    pdf_meta.update(ext_meta)

    parsed = ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type="pdf",
        metadata=pdf_meta,
    )

    text_parts = []
    md_parts = [f"# {parsed.filename}\n"]

    for i, page in enumerate(pdf_doc, 1):
        page_text = page.get_text()
        if page_text.strip():
            text_parts.append(page_text)
            md_parts.append(f"\n## 第 {i} 页\n")
            md_parts.append(page_text)

    pdf_doc.close()
    parsed.text_content = "\n".join(text_parts)
    parsed.structured_content = "\n".join(md_parts)
    parsed.summary = f"共 {ext_meta['page_count']} 页"
    return parsed


# ── .xlsx 解析 ──────────────────────────────────────────


def parse_xlsx(filepath: str) -> ParsedDocument:
    import openpyxl

    parsed = ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type="xlsx",
        metadata=_build_metadata(filepath),
    )

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    parsed.metadata["sheet_count"] = len(sheet_names)

    text_parts = []
    md_parts = [f"# {parsed.filename}\n"]
    total_rows = 0

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        md_parts.append(f"\n## Sheet: {sheet_name}\n")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            if row_count > 10000:
                break
            cells = [str(c) if c is not None else "" for c in row]
            text_parts.append(" | ".join(cells))
            md_parts.append("| " + " | ".join(cells) + " |")
        total_rows += row_count
        parsed.metadata[f"sheet_{sheet_name}_rows"] = row_count

    wb.close()
    parsed.text_content = "\n".join(text_parts)
    parsed.structured_content = "\n".join(md_parts)
    parsed.summary = f"共 {len(sheet_names)} 个Sheet, {total_rows} 行数据"
    return parsed


# ── .xls 解析（pandas calamine 引擎） ──────────────────


def parse_xls(filepath: str) -> ParsedDocument:
    import pandas as pd

    parsed = ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type="xls",
        metadata=_build_metadata(filepath),
    )

    dfs = pd.read_excel(filepath, sheet_name=None, engine="calamine")
    parsed.metadata["sheet_count"] = len(dfs)

    text_parts = []
    md_parts = [f"# {parsed.filename}\n"]
    total_rows = 0

    for sheet_name, df in dfs.items():
        md_parts.append(f"\n## Sheet: {sheet_name}\n")
        row_count = 0
        for _, row in df.head(10000).iterrows():
            row_count += 1
            cells = [str(v) if pd.notna(v) else "" for v in row]
            text_parts.append(" | ".join(cells))
            md_parts.append("| " + " | ".join(cells) + " |")
        total_rows += row_count
        parsed.metadata[f"sheet_{sheet_name}_rows"] = row_count

    parsed.text_content = "\n".join(text_parts)
    parsed.structured_content = "\n".join(md_parts)
    parsed.summary = f"共 {len(dfs)} 个Sheet, {total_rows} 行数据"
    return parsed


# ── .pptx 解析 ──────────────────────────────────────────


def parse_pptx(filepath: str) -> ParsedDocument:
    from pptx import Presentation

    parsed = ParsedDocument(
        filepath=filepath,
        filename=os.path.basename(filepath),
        file_type="pptx",
        metadata=_build_metadata(filepath),
    )

    prs = Presentation(filepath)
    parsed.metadata["slide_count"] = len(prs.slides)

    text_parts = []
    md_parts = [f"# {parsed.filename}\n"]

    for i, slide in enumerate(prs.slides, 1):
        md_parts.append(f"\n## 幻灯片 {i}\n")
        slide_texts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_texts.append(t)

            if shape.has_table:
                table = shape.table
                table_texts = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_texts.append(" | ".join(cells))
                slide_texts.extend(table_texts)

        text_parts.extend(slide_texts)
        md_parts.append("\n".join(slide_texts))

    parsed.text_content = "\n".join(text_parts)
    parsed.structured_content = "\n".join(md_parts)
    parsed.summary = f"共 {len(prs.slides)} 张幻灯片"
    return parsed


# ── 辅助函数 ────────────────────────────────────────────


def _summarize_text(text_parts: list[str]) -> str:
    total = len(text_parts)
    first_line = text_parts[0][:80] if text_parts else ""
    return f"共 {total} 段。首段：{first_line}..."